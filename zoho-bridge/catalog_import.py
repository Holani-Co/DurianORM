# Product price-list import for Settings → Product Catalog. Turns what the client
# provides into the catalog document product_catalog.py serves:
#   {"price_period", "source", "source_name", "products": {sku: {...}}}
#
# Two accepted shapes, detected automatically:
#   - SIMPLE TEMPLATE (a Google Sheet, .xlsx or .csv): one tab, a header row with
#     SKU | Name | MRP | Sale price | Category (order and case don't matter;
#     common aliases accepted). Made for hand-editing; export_csv() writes it.
#   - MONTHLY WORKBOOK: the pricing team's file with 'Durian' + 'RT' tabs
#     (header on row 2, data from row 3) — merged exactly as the original offline
#     build did (Durian.Offer price wins, RT revised sale price as fallback).
#
# A linked Google Sheet is read through Google's own export link, so it must be
# shared "Anyone with the link can view". Only docs.google.com spreadsheet URLs
# are fetched — the bridge never follows an arbitrary URL.

import csv
import io
import re
from datetime import datetime, timezone

import httpx
import openpyxl

MAX_FILE_BYTES = 15 * 1024 * 1024
# More than this share of rows unusable (and at least MIN_BAD_TO_REFUSE of them)
# → refuse: it is almost certainly the wrong file or a shifted header, not a few
# bad rows. The floor keeps one blank row in a small hand-made sheet publishable.
MAX_BAD_SHARE = 0.2
MIN_BAD_TO_REFUSE = 5

_COLUMNS = {
    "sku": ("sku", "sku id", "sku code", "code", "item code"),
    "name": ("name", "product name", "description", "product"),
    "mrp": ("mrp",),
    "sale_price": ("sale price", "selling price", "offer price", "price", "sale"),
    "category": ("category",),
    "family": ("family",),
}


class CatalogImportError(Exception):
    """A whole-file problem shown to the admin (wrong file, sheet not shared…)."""


# ── cell helpers (shared with the monthly-workbook merge) ─────────────────────

def _num(v):
    try:
        f = float(str(v).replace(",", "").replace("₹", "").strip())
        return f if f > 0 else None
    except (TypeError, ValueError):
        return None


def _sku(v):
    s = str(v or "").strip().upper()
    return s if s and s not in ("#N/A", "NONE") else None


def _clean(v):
    s = str(v or "").strip()
    return "" if s in ("", "#N/A", "None") else s


def _head(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip().lower())


# ── monthly workbook ('Durian' + 'RT' tabs) ───────────────────────────────────

def is_monthly_workbook(wb) -> bool:
    return "Durian" in wb.sheetnames and "RT" in wb.sheetnames


def parse_monthly_workbook(wb) -> dict:
    """{products, stats} from the pricing team's workbook. Column positions are
    the verified schema of that file (see build_product_catalog.py)."""
    durian = {}
    for row in wb["Durian"].iter_rows(min_row=3, values_only=True):
        row = (tuple(row) + (None,) * 9)[:9]
        sku = _sku(row[0])
        if sku:
            durian[sku] = {"mrp": _num(row[1]), "offer": _num(row[3]),
                           "category": _clean(row[8])}
    rt = {}
    for row in wb["RT"].iter_rows(min_row=3, values_only=True):
        row = (tuple(row) + (None,) * 10)[:10]
        sku = _sku(row[0])
        if sku:
            rt[sku] = {"name": _clean(row[1]), "mrp": _num(row[3]),
                       "revised_sale": _num(row[7]), "category": _clean(row[9])}

    products, mismatches, unpriced = {}, 0, []
    for sku in sorted(set(durian) | set(rt)):
        d, t = durian.get(sku, {}), rt.get(sku, {})
        offer, revised = d.get("offer"), t.get("revised_sale")
        if offer and revised and abs(offer - revised) / max(offer, revised) > 0.01:
            mismatches += 1
        sale = offer or revised
        if not sale:
            unpriced.append(sku)          # no sale price → can't quote EMI/price
            continue
        products[sku] = {
            "name": t.get("name") or "",
            "family": sku.split("/")[0],
            "mrp": d.get("mrp") or t.get("mrp"),
            "sale_price": round(sale),
            "category": d.get("category") or t.get("category") or "",
            "source": "durian" if offer else "rt",
        }
    warnings = []
    if unpriced:
        warnings.append(f"{len(unpriced)} SKUs skipped — no sale price "
                        f"(e.g. {', '.join(unpriced[:5])})")
    if mismatches:
        warnings.append(f"{mismatches} SKUs where the Durian offer price and the RT "
                        f"revised price differ by more than 1% (Durian price used)")
    return {"products": products, "warnings": warnings,
            "stats": {"durian_rows": len(durian), "rt_rows": len(rt),
                      "mismatches": mismatches}}


# ── simple template ───────────────────────────────────────────────────────────

def _find_header(rows: list[tuple]) -> tuple[int, dict] | tuple[None, None]:
    """(header row index, {field: column}) from the first rows, or (None, None)."""
    for i, row in enumerate(rows[:5]):
        heads = [_head(c) for c in row]
        cols = {}
        for field, aliases in _COLUMNS.items():
            for alias in aliases:
                if alias in heads:
                    cols[field] = heads.index(alias)
                    break
        if "sku" in cols and "sale_price" in cols:
            return i, cols
    return None, None


def parse_simple_rows(rows: list[tuple]) -> dict:
    """{products, warnings} from a header + rows table (CSV or one sheet tab)."""
    hi, cols = _find_header(rows)
    if hi is None:
        raise CatalogImportError("Couldn't find the header row. The first rows need "
                                 "columns named SKU and Sale price (plus Name, MRP, "
                                 "Category).")

    def cell(row, field):
        c = cols.get(field)
        return row[c] if c is not None and c < len(row) else None

    products, bad, dupes, data_rows = {}, [], [], 0
    for n, row in enumerate(rows[hi + 1:], start=hi + 2):
        if not any(_clean(c) for c in row):
            continue
        data_rows += 1
        sku = _sku(cell(row, "sku"))
        sale = _num(cell(row, "sale_price"))
        if not sku or not sale:
            bad.append(f"row {n}: " + ("no SKU" if not sku else f"{sku} has no sale price"))
            continue
        if sku in products:
            dupes.append(sku)
        products[sku] = {
            "name": _clean(cell(row, "name")),
            "family": _clean(cell(row, "family")).upper() or sku.split("/")[0],
            "mrp": _num(cell(row, "mrp")),
            "sale_price": round(sale),
            "category": _clean(cell(row, "category")),
            "source": "sheet",
        }
    if not products:
        raise CatalogImportError("No usable products found — every row is missing a "
                                 "SKU or a sale price.")
    if len(bad) >= MIN_BAD_TO_REFUSE and len(bad) / data_rows > MAX_BAD_SHARE:
        raise CatalogImportError(f"{len(bad)} of {data_rows} rows are missing a SKU or "
                                 f"sale price — check this is the right sheet and the "
                                 f"header row (first: {'; '.join(bad[:3])}).")
    warnings = []
    if bad:
        warnings.append(f"{len(bad)} rows skipped: " + "; ".join(bad[:10])
                        + (" …" if len(bad) > 10 else ""))
    if dupes:
        warnings.append(f"{len(dupes)} duplicate SKUs — the last row wins "
                        f"(e.g. {', '.join(sorted(set(dupes))[:5])})")
    return {"products": products, "warnings": warnings}


# ── entry points ──────────────────────────────────────────────────────────────

def parse_file(data: bytes, filename: str = "") -> dict:
    """Parse an uploaded/fetched file. Returns {products, warnings, format}."""
    if not data:
        raise CatalogImportError("The file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise CatalogImportError("The file is larger than 15 MB.")
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        res = parse_simple_rows([tuple(r) for r in csv.reader(io.StringIO(text))])
        return {**res, "format": "template"}
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise CatalogImportError("Couldn't read the file — upload an Excel (.xlsx) or "
                                 ".csv file.")
    try:
        if is_monthly_workbook(wb):
            res = parse_monthly_workbook(wb)
            if not res["products"]:
                raise CatalogImportError("The workbook's Durian/RT tabs have no priced rows.")
            return {"products": res["products"], "warnings": res["warnings"],
                    "format": "workbook"}
        rows = [tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        return {**parse_simple_rows(rows), "format": "template"}
    finally:
        wb.close()


_SHEET_RE = re.compile(r"^https://docs\.google\.com/spreadsheets/d/([A-Za-z0-9_-]{20,})")


def sheet_id(url: str) -> str | None:
    m = _SHEET_RE.match(str(url or "").strip())
    return m.group(1) if m else None


async def fetch_sheet(url: str) -> bytes:
    """The whole Google Sheet as .xlsx (so both accepted shapes work)."""
    sid = sheet_id(url)
    if not sid:
        raise CatalogImportError("That isn't a Google Sheets link — it should start "
                                 "with https://docs.google.com/spreadsheets/d/…")
    export = f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx"
    try:
        async with httpx.AsyncClient(timeout=30, follow_redirects=True) as client:
            r = await client.get(export)
    except httpx.HTTPError as e:
        raise CatalogImportError(f"Couldn't reach Google Sheets ({type(e).__name__}).")
    ctype = r.headers.get("content-type", "")
    if r.status_code in (401, 403) or "text/html" in ctype:
        raise CatalogImportError("Google refused access. Share the sheet as \"Anyone "
                                 "with the link can view\" and try again.")
    if r.status_code == 404:
        raise CatalogImportError("Google Sheets says this sheet doesn't exist.")
    if r.status_code >= 300:
        raise CatalogImportError(f"Google Sheets returned an error ({r.status_code}).")
    return r.content


def build_doc(parsed: dict, source: str, source_name: str = "",
              price_period: str = "") -> dict:
    """The catalog document to publish."""
    return {"products": parsed["products"], "source": source,
            "source_name": source_name, "format": parsed.get("format", ""),
            "price_period": price_period.strip()
            or f"Updated {datetime.now(timezone.utc).strftime('%d %b %Y')}"}


def diff(old: dict, new: dict, sample: int = 15) -> dict:
    """What publishing `new` products changes against the live `old` products."""
    old, new = old or {}, new or {}
    added = sorted(set(new) - set(old))
    removed = sorted(set(old) - set(new))
    changed, details = [], 0
    for sku in sorted(set(old) & set(new)):
        a, b = old[sku].get("sale_price"), new[sku].get("sale_price")
        if a != b:
            changed.append({"sku": sku, "name": new[sku].get("name") or "",
                            "old": a, "new": b})
        elif any(old[sku].get(k) != new[sku].get(k)
                 for k in ("name", "mrp", "category", "family")):
            details += 1

    def rows(skus, src):
        return [{"sku": s, "name": src[s].get("name") or "",
                 "sale_price": src[s].get("sale_price")} for s in skus[:sample]]
    return {"count": len(new), "previous_count": len(old),
            "added": len(added), "removed": len(removed),
            "price_changed": len(changed), "details_changed": details,
            "unchanged": len(new) - len(added) - len(changed) - details,
            "added_sample": rows(added, new), "removed_sample": rows(removed, old),
            "changed_sample": changed[:sample]}


def export_csv(products: dict) -> str:
    """The current list in the simple-template shape — the starting point for a
    Google Sheet or a re-upload."""
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["SKU", "Name", "MRP", "Sale price", "Category", "Family"])
    for sku in sorted(products):
        p = products[sku]
        mrp = p.get("mrp")
        w.writerow([sku, p.get("name") or "",
                    int(mrp) if isinstance(mrp, (int, float)) else "",
                    p.get("sale_price") or "", p.get("category") or "",
                    p.get("family") or ""])
    return out.getvalue()
