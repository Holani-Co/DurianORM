# Build the unified store registry (data/store_registry.json) — the FLOOR that
# store_locator reads (Settings → Stores overrides merge on top at runtime).
#
#   Pass 1 — the client's store-locator export is the ROUTING spine: id, lat/lon,
#            pincode, ZOHO owner and verticals (a store's own "Verticle" plus any
#            vertical in its "Service" column). Active stores only. FHC owners,
#            blank in the export, come from routing_rules.yaml.
#   Pass 2 — the "Master ORM Store List" supplies the current DISPLAY details
#            (address, timing + closed day, phone, manager, email, real map link,
#            holiday, parking, store type, area, floors, escalator). Nothing in the
#            sheet goes live unless the REVIEWED map data/master_store_map.json
#            names it:
#              mapping      sheet name → registry id: display overlay (the store's
#                           own name, city and routing fields are kept)
#              new_stores   sheet name → {card_name, owner_id?, lat?, lon?}
#              coming_soon  sheet name → registry id to take off the floor (or null)
#              card_names   registry id → reviewed display-name correction
#            A sheet name the map doesn't know is reported and skipped.
#
# Re-run whenever the client sends a fresh export or Master list.

import json
import re
from pathlib import Path

import openpyxl
import yaml

_HERE = Path(__file__).parent
_XLSX = _HERE / "data" / "store_locator_export2026_03_18_11_22_08_242251 (1).xlsx"
_MASTER = _HERE / "data" / "Master ORM Store List.xlsx"
_MASTER_MAP = _HERE / "data" / "master_store_map.json"
_OUT = _HERE / "data" / "store_registry.json"
_ROUTING = _HERE / "routing_rules.yaml"

# Sheet label → internal vertical token used across the codebase.
_VERT = {"furniture": "furniture", "door": "doors", "doors": "doors",
         "full home customisation": "fhc", "fhc": "fhc"}

# Master ORM list columns, BY POSITION — the sheet has DUPLICATE headers (two
# "Vertical" at 2 & 5, two "Week Off" at 18 & 20), so hdr.index() would grab the
# wrong column. Positions are stable across the client's exports.
_M = {"vert": 2, "service": 3, "name": 4, "store_type": 6, "area": 7,
      "floors": 8, "addr": 9, "city": 10, "timing": 11, "phone": 12,
      "email": 13, "holiday": 14, "map": 15, "manager": 17, "parking": 21,
      "escalator": 22}

# A sheet cell holding one of these means "no value".
_PLACEHOLDERS = {"-", "na", "n/a", "nil", "none", "null"}
_DAYS = ("monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday")


def _vert_tokens(*cells) -> list[str]:
    """Normalise a Verticle cell and/or a Service cell into internal tokens."""
    out: list[str] = []
    for cell in cells:
        for part in re.split(r"[,/;]", str(cell or "")):
            key = part.strip().lower()
            tok = _VERT.get(key)
            if tok and tok not in out:
                out.append(tok)
    return out


def _slug(name: str, fallback: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", str(name or "").lower()).strip("-")
    return s or fallback


def _map_url(lat, lon):
    if lat in (None, "") or lon in (None, ""):
        return ""
    return f"https://www.google.com/maps/search/?api=1&query={lat},{lon}"


def _load_fhc_owners() -> dict:
    """Studio location-fragment → owner_id, from crm_owner_routing_homestudio.
    Keyed by a distinctive fragment ('kirti nagar') so we can match the sheet's
    longer store names."""
    try:
        rules = yaml.safe_load(_ROUTING.read_text()) or {}
    except Exception as e:
        print(f"[registry] could not read routing_rules.yaml: {e}")
        return {}
    homestudio = rules.get("crm_owner_routing_homestudio") or {}
    owners = {}
    for loc, info in homestudio.items():
        oid = str((info or {}).get("owner_id") or "").strip()
        if not oid:
            continue
        frag = loc.split("-")[-1].strip().lower()   # "Delhi - Kirti Nagar" → "kirti nagar"
        owners[frag] = oid
    return owners


def _text(v) -> str:
    """A cell value as one clean line: whole-number floats ("14500.0") as
    integers, whitespace runs collapsed, placeholders ("-", "NA", "None") as ""."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = " ".join(str(v if v is not None else "").split())
    return "" if t.lower() in _PLACEHOLDERS else t


def _mraw(r, key: str) -> str:
    """A Master cell as text: whole-number floats ("14500.0") as integers and
    placeholders ("-", "NA") as "". Line breaks are kept (Manager Details needs
    them) — use _mcell for single-line values."""
    i = _M[key]
    v = r[i] if i < len(r) else None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = str(v).strip() if v is not None else ""
    return "" if t.lower() in _PLACEHOLDERS else t


def _mcell(r, key: str) -> str:
    """A Master cell by position, as one clean line (see _text)."""
    i = _M[key]
    return _text(r[i] if i < len(r) else None)


def _master_manager(r) -> str:
    # "Name\nphone[\nemail]", sometimes all on one line ("MANJARI SARKAR -
    # 9870570655") — the name is what precedes the first digit or e-mail.
    first = next((ln for ln in _mraw(r, "manager").splitlines() if ln.strip()), "")
    name = " ".join(re.split(r"[\d@]", first)[0].split()).strip(" -:,")
    return "" if name.lower() in _PLACEHOLDERS else name


def _master_phone(r) -> str:
    # Just the numbers ("Mr. Monish: 8618857149/ 7715966829", a leading emoji).
    t = re.sub(r"[^\d/+\- ]", "", _mcell(r, "phone"))
    t = re.sub(r"\s*/\s*", " / ", " ".join(t.split())).strip(" -/")
    return t if re.search(r"\d{6}", t.replace(" ", "")) else ""


def _closed_day(holiday: str) -> str:
    """"Sunday" / "Closed On: Tuesday" → the weekday; "No Holiday" → ""."""
    day = re.sub(r"(?i)closed\s*on\s*:?", "", holiday).strip().lower()
    return day.title() if day in _DAYS else ""


def _extract_pincode(addr: str) -> str:
    # The LAST 6-digit run: an address's pincode trails it, and a plot number can
    # come first (pincode_resolver.extract_pincode takes the first — for DMs).
    hits = re.findall(r"(?<!\d)([1-9]\d{5})(?!\d)", addr or "")
    return hits[-1] if hits else ""


def _master_display(r) -> dict:
    """The display fields one Master row supplies (non-empty only, so an overlay
    never blanks a good registry value)."""
    timing, closed = _mcell(r, "timing"), _closed_day(_mcell(r, "holiday"))
    if timing and closed and "closed" not in timing.lower():
        timing = f"{timing}, Closed on {closed}"    # the card prints timing, not holiday
    link = (_mcell(r, "map").split() or [""])[0]
    d = {
        "address": _mcell(r, "addr"),
        "timing": timing,
        "phone": _master_phone(r),
        "email": _mcell(r, "email"),
        "manager": _master_manager(r),
        "map_url": link if link.startswith("http") else "",
        "holiday": _mcell(r, "holiday"),
        "parking": _mcell(r, "parking"),
        "store_type": _mcell(r, "store_type").lower(),
        "area": _mcell(r, "area"),
        "floors": _mcell(r, "floors"),
        "escalator": _mcell(r, "escalator"),
    }
    return {k: v for k, v in d.items() if v}


def _new_store(r, name: str, disp: dict, review: dict, stores: list) -> dict:
    """A store that exists only in the Master list: its row plus the reviewed map
    entry (card name, CRM owner, optional lat/lon). Coordinates come from the
    review or an existing store at the same pincode — never a pincode centroid
    (the geocoder puts most of a metro on one point, which would make the new
    store "nearest" for the whole city). Without coordinates it still matches its
    exact pincode and its name."""
    verts = _vert_tokens(_mcell(r, "vert"), _mcell(r, "service"))
    pin = _extract_pincode(disp.get("address", ""))
    twin = next((s for s in stores if pin and s.get("pincode") == pin
                 and s.get("lat") is not None), None)
    sid = _slug(name, name)
    if any(s["id"] == sid for s in stores):
        raise SystemExit(f"[registry] new store '{name}' collides with id '{sid}'")
    return {
        "id": sid, "name": name, "card_name": review["card_name"],
        "verticals": verts,
        "primary_vertical": (_vert_tokens(_mcell(r, "vert")) or verts)[0],
        "city": _mcell(r, "city"), "location": "", "state": "", "facility": "",
        "image_url": "", "pincode": pin,
        "lat": review.get("lat", twin["lat"] if twin else None),
        "lon": review.get("lon", twin["lon"] if twin else None),
        "crm_owner_id": review.get("owner_id", ""),
        **{k: disp.get(k, "") for k in ("address", "timing", "phone", "email",
                                         "manager", "map_url", "holiday", "parking",
                                         "store_type", "area", "floors", "escalator")},
    }


def _merge_master(stores: list[dict]) -> list[dict]:
    """Pass 2 (see the header). A missing sheet or map fails the build rather
    than silently writing a registry without the Master details."""
    m = json.loads(_MASTER_MAP.read_text())
    mapping, new_stores, coming_soon = m["mapping"], m["new_stores"], m["coming_soon"]
    ws = openpyxl.load_workbook(_MASTER, data_only=True)["Master ORM Store List"]
    by_id = {s["id"]: s for s in stores}
    seen, added, unknown, overlaid = set(), [], [], 0
    for r in list(ws.iter_rows(values_only=True))[1:]:
        name = _mcell(r, "name")
        if not name or name in coming_soon:            # blank rows; dropped below
            seen.add(name)
            continue
        seen.add(name)
        disp = _master_display(r)
        if name in mapping:
            store = by_id[mapping[name]]
            # Keep the registry's number when it already holds this one — the
            # sheet's numeric cell drops a landline's leading 0 (02402350686).
            if disp.get("phone") and re.sub(r"\D", "", disp["phone"]) in \
                    re.sub(r"\D", "", store.get("phone") or ""):
                del disp["phone"]
            store.update(disp)
            overlaid += 1
        elif name in new_stores:
            added.append(_new_store(r, name, disp, new_stores[name], stores))
        else:
            unknown.append(name)
    for rid, card in m.get("card_names", {}).items():    # reviewed name fixes
        by_id[rid]["card_name"] = card
    drop = {rid for rid in coming_soon.values() if rid}
    stores = [s for s in stores if s["id"] not in drop] + added
    print(f"[registry] master: {overlaid} display-updated, {len(added)} new, "
          f"{len(drop)} coming-soon taken off the floor")
    for n in unknown:
        print(f"[registry] WARNING: '{n}' is in the sheet but not in "
              f"master_store_map.json — skipped; review it and add it to the map")
    for n in sorted((set(mapping) | set(new_stores) | set(coming_soon)) - seen):
        print(f"[registry] WARNING: the map names '{n}', which is not in the sheet")
    return stores


def build() -> list[dict]:
    wb = openpyxl.load_workbook(_XLSX, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() if h is not None else "" for h in rows[0]]

    def idx(name):
        return hdr.index(name)

    cN, cPh, cAddr, cCity, cLoc, cState, cPin, cLat, cLon, cMgr, cEmail, \
        cFac, cTime, cType, cStreet, cSlug, cVert, cImg2, cOwner, cSvc, cStatus = (
            idx("Name"), idx("Phone Number"), idx("Adderss"), idx("City"),
            idx("Location"), idx("State"), idx("Pincode"), idx("Latitude"),
            idx("Longitude"), idx("Manager Name"), idx("Store Email"),
            idx("Store Facility"), idx("Store Timing"), idx("Store Type"),
            idx("Street View"), idx("Slug"), idx("Verticle"),
            idx("Store Image Link"), idx("ZOHO Owner ID"), idx("Service"),
            idx("Status"))

    fhc_owners = _load_fhc_owners()
    out, skipped_inactive, fhc_merged, no_owner = [], 0, 0, 0

    for n, r in enumerate(rows[1:], start=2):
        if not any(c not in (None, "") for c in r):
            continue
        if str(r[cStatus]).strip().lower() != "active":
            skipped_inactive += 1
            continue
        verticals = _vert_tokens(r[cVert], r[cSvc])
        if not verticals:
            continue
        name = str(r[cN]).strip()
        try:
            lat = float(r[cLat]); lon = float(r[cLon])
        except (TypeError, ValueError):
            lat = lon = None
        owner = str(r[cOwner]).strip() if r[cOwner] not in (None, "") else ""
        # FHC rows carry no owner in this export → merge from routing_rules.
        if not owner and "fhc" in verticals:
            low = name.lower()
            for frag, oid in fhc_owners.items():
                if frag and frag in low:
                    owner = oid
                    fhc_merged += 1
                    break
        if not owner:
            no_owner += 1
        out.append({
            "id": _slug(r[cSlug] or name, f"store-{n}"),
            "name": name,
            "card_name": _text(name.replace("Durian ", "")),
            "verticals": verticals,
            "primary_vertical": _vert_tokens(r[cVert])[0] if _vert_tokens(r[cVert]) else verticals[0],
            "store_type": str(r[cType] or "").strip().lower(),
            "city": _text(r[cCity]),
            "location": _text(r[cLoc]),
            "state": _text(r[cState]),
            "address": _text(r[cAddr]),
            "pincode": re.sub(r"\D", "", str(r[cPin] or "")) or "",
            "lat": lat, "lon": lon,
            "manager": _text(r[cMgr]),
            "phone": _text(r[cPh]),
            "email": _text(r[cEmail]),
            "timing": _text(r[cTime]),
            "facility": _text(r[cFac]),
            "map_url": (str(r[cStreet]).strip() if r[cStreet] not in (None, "")
                        else _map_url(lat, lon)),
            "image_url": str(r[cImg2] or "").strip(),
            "crm_owner_id": owner,
        })

    print(f"[registry] export spine: {len(out)} active stores "
          f"(skipped {skipped_inactive} inactive; FHC owners merged {fhc_merged}; "
          f"no owner id {no_owner})")
    out = _merge_master(out)
    _OUT.write_text(json.dumps(out, ensure_ascii=False, indent=2))
    print(f"[registry] wrote {len(out)} stores → {_OUT.name}")
    return out


if __name__ == "__main__":
    build()
